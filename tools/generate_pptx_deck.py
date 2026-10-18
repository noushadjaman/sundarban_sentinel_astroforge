import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    slides_config = [
        {
            "num": 1,
            "title": "Title & Team AstroForge",
            "image": "project_assets/hackathon_slides/slide_1_title_astroforge.png",
            "notes": (
                "SLIDE 1 TALKING POINTS (TEAM ASTROFORGE):\n"
                "- Welcome honorable judges and participants of Hack For Humanity 2026. We are Team AstroForge, represented by Md Nushad Jaman Raj, Tanjid Hossen Rifat, and Raf Rafin Mahmud.\n"
                "- Today we present Sundarban Sentinel: a fully engineered, zero-mock production distributed system built for mangrove conservation and frontline humanitarian protection.\n"
                "- The Sundarbans is the world's largest contiguous mangrove forest (10,000 km²) and a UNESCO World Heritage site that protects over 30 million coastal citizens from deadly cyclone storm surges.\n"
                "- However, the deep forest is a total cellular black hole. Forest rangers, honey collectors (Mawalis), and fishermen operate in complete isolation with 0 bars of mobile coverage.\n"
                "- Our solution brings on-device edge AI, decoupled statutory legal deterrence, live NASA Earthdata satellite telemetry, and auto-syncing distributed cloud infrastructure to protect this critical ecosystem."
            )
        },
        {
            "num": 2,
            "title": "The Crisis: Poaching, Poison Fishing & Isolation",
            "image": "project_assets/hackathon_slides/slide_news_crisis_overview.png",
            "notes": (
                "SLIDE 2 TALKING POINTS (PROBLEM STATEMENT & GROUND EVIDENCE):\n"
                "- Here is the undeniable ground reality in the Sundarbans, documented by verified investigative reports in The Business Standard and Dhaka Tribune.\n"
                "- Crisis 1: Poison Fishing Syndicates pour toxic pesticide chemicals into tidal canals to force instant fish harvests, causing catastrophic aquatic die-offs and poisoning food chains.\n"
                "- Crisis 2: Lethal Current Nets & Snares. Tens of thousands of meters of banned monofilament nylon current jal trap endangered Ganges river dolphins and Olive Ridley turtles.\n"
                "- Crisis 3: Apex Predator Threat & Venison Syndicates. Wire snares kill adult Royal Bengal Tigers, while poachers smuggle hundreds of kilograms of spotted deer venison (824kg seized in single raids).\n"
                "- Crisis 4: The 10,000 km² Communication Void. Rangers patrol blind tidal channels for up to 14 days without cellular coverage, rendering conventional cloud apps completely useless."
            )
        },
        {
            "num": 3,
            "title": "Topic 1: 100% Offline Terminal Mode",
            "image": "project_assets/hackathon_slides/slide_topic_1_offline_terminal.png",
            "notes": (
                "SLIDE 3 TALKING POINTS (TOPIC 1: 100% OFFLINE RESILIENCE):\n"
                "- Topic 1 demonstrates our core innovation: 100% autonomous offline terminal mode.\n"
                "- In FIG 1.1 and 1.2, you can see our live Flutter application running with ZERO cellular bars.\n"
                "- Point 1: Cached NASA Anomaly Layer. Satellite risk polygons are pre-cached on device flash memory.\n"
                "- Point 2: Offline GPS River Tracking. Real-time river coordinate lock without cell tower triangulation.\n"
                "- Point 3: On-Device INT8 YOLOv8 Model. Neural weights execute entirely on the phone's CPU/NPU with 353ms inference latency.\n"
                "- Point 4: Zero-Data-Loss SQLite Vault. Incidents and GPS coordinates are buffered securely in an encrypted local queue until network return."
            )
        },
        {
            "num": 4,
            "title": "Topic 2: AI Species Legal Verification for Hunters",
            "image": "project_assets/hackathon_slides/slide_topic_2_edge_ai_species.png",
            "notes": (
                "SLIDE 4 TALKING POINTS (TOPIC 2: HUNTER LEGAL GUIDANCE & AI DETECTION):\n"
                "- Topic 2 addresses professional trackers, hunters, and Mawalis (traditional honey gatherers).\n"
                "- In FIG 2.1 and 2.2, rangers or hunters simply point their mobile camera at wildlife.\n"
                "- Our quantized neural network classifies the Royal Bengal Tiger at 94% confidence and Spotted Deer at 92% confidence in 353ms.\n"
                "- Crucially, our decoupled statutory engine maps the detection directly to the Bangladesh Wildlife Conservation Act 2012 (Section 36).\n"
                "- It instantly generates bilingual statutory alert cards stating non-bailable status, 2 to 7 years rigorous imprisonment, and up to 10 lakh taka fines.\n"
                "- This provides clear deterrence, educating coastal communities before any shot, trap, or harm occurs."
            )
        },
        {
            "num": 5,
            "title": "Topic 3: Community Net Checks & Citizen Rescue",
            "image": "project_assets/hackathon_slides/slide_topic_3_general_people_usecase.png",
            "notes": (
                "SLIDE 5 TALKING POINTS (TOPIC 3: CITIZEN EMPOWERMENT & NET CHECKS):\n"
                "- In Topic 3, we empower general citizens and coastal fishermen across 4 daily use cases:\n"
                "- Use Case 1: Pre-Fishing Net Check. Fishermen scan their fishing nets on the riverbank before casting to ensure mesh legality under the Protection and Conservation of Fish Act 1950, avoiding boat confiscation.\n"
                "- Use Case 2: Legal Awareness without Jargon. Plain-language statutory guidance in Bengali.\n"
                "- Use Case 3: Endangered Marine Fauna Protection. Preventing the accidental catch of Ganges River Dolphins and Olive Ridley turtles.\n"
                "- Use Case 4: Citizen Deforestation & Animal Rescue Desk. Villagers capture photos of injured animals trapped in mud or illegal tree-felling, queuing an instant GPS dossier for rangers."
            )
        },
        {
            "num": 6,
            "title": "Topic 4: NASA Satellite Deforestation Radar",
            "image": "project_assets/hackathon_slides/slide_topic_4_nasa_deforestation.png",
            "notes": (
                "SLIDE 6 TALKING POINTS (TOPIC 4: NASA SATELLITE RADAR FOR RANGERS):\n"
                "- Topic 4 introduces our remote sensing engine for forest department rangers.\n"
                "- Our backend pipeline queries NASA Earthdata CMR endpoints for MODIS MOD13Q1 (250m NDVI) and MOD11A2 (surface temperature).\n"
                "- Our anomaly algorithm identifies ΔNDVI drops exceeding -0.15, signaling active illegal timber felling camps.\n"
                "- In FIG 4.1 and 4.2, the mobile radar shows high-risk canopy drop polygons directly on the ranger's offline map.\n"
                "- The one-tap compass calculates the exact steering bearing through tidal channels, directing patrol speedboats directly to poacher camps without fuel waste."
            )
        },
        {
            "num": 7,
            "title": "Topic 5: Cellular Handshake & Auto-Sync",
            "image": "project_assets/hackathon_slides/slide_topic_5_cellular_sync.png",
            "notes": (
                "SLIDE 7 TALKING POINTS (TOPIC 5: TELECOM HANDSHAKE & CLOUD INGESTION):\n"
                "- Topic 5 demonstrates what happens when patrol boats or villagers return to coastal cell coverage.\n"
                "- An automated background handshake connects to coastal 4G/5G towers with zero user intervention.\n"
                "- Point 1: Two-way sync fetches updated NASA thermal anomaly polygons.\n"
                "- Point 2: Our background daemon uploads all local SQLite incidents and citizen reports directly to our Cloudflare Workers edge gateway.\n"
                "- Point 3: Media photos are uploaded directly to our Cloudinary CDN preset, avoiding server load and committing evidence directly into Firestore."
            )
        },
        {
            "num": 8,
            "title": "Production Architecture & Team AstroForge",
            "image": "project_assets/hackathon_slides/slide_8_architecture_team.png",
            "notes": (
                "SLIDE 8 TALKING POINTS (PRODUCTION SYNTHESIS & CLOSING):\n"
                "- In conclusion, Sundarban Sentinel is not a mockup or a theoretical pitch.\n"
                "- Every single tier—from live NASA Earthdata CMR ingestion to on-device INT8 YOLOv8 inference and Cloudflare Workers gateway—is fully functional and production-tested.\n"
                "- On behalf of Team AstroForge:\n"
                "  • Md Nushad Jaman Raj (Team Lead & Mobile Edge AI)\n"
                "  • Tanjid Hossen Rifat (Cloud & Edge Gateway Architect)\n"
                "  • Raf Rafin Mahmud (GIS & Remote Sensing Engineer)\n"
                "- We believe AI must serve the planet's most vulnerable ecosystems and coastal communities.\n"
                "- Thank you, Hack For Humanity judges and organizers. We are now open for questions!"
            )
        }
    ]

    for item in slides_config:
        slide = prs.slides.add_slide(blank_layout)
        
        # Ensure base background is light
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = RGBColor(248, 250, 252)
        bg_shape.line.fill.background()

        # Add full-bleed 1080p slide graphic (13.333" x 7.5")
        img_path = os.path.abspath(item["image"])
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, 0, 0, prs.slide_width, prs.slide_height)
        else:
            print(f"Warning: Image not found: {img_path}")

        # Add Speaker Notes
        slide.notes_slide.notes_text_frame.text = item["notes"]

    out_file = os.path.abspath('Sundarban_Sentinel_Hack_For_Humanity_Pitch_Deck.pptx')
    prs.save(out_file)
    print(f"SUCCESS: Generated {len(prs.slides)} slides in {out_file}")

if __name__ == '__main__':
    build_deck()
